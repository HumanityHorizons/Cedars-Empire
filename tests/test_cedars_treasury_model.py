from io import BytesIO
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook

from cedars_treasury_model import (
    DEFAULTS,
    _wb_bytes,
    capped_distribution,
    fee_preview,
    solvency_status,
)


def test_distribution_requires_approval_and_caps_each_investor():
    assert capped_distribution(10000, 20000, 1000, "0.01", approved=False) == 0
    one = capped_distribution(10000, 20000, 30, "0.01", approved=True)
    two = capped_distribution(10000, 20000, 30, "0.01", approved=True)
    assert one == 15
    assert two == 15
    assert one + two <= 30
    assert capped_distribution(10000, 20000, 10000, "0.01", approved=True) <= 100


def test_distribution_rejects_invalid_or_ineligible_inputs():
    assert capped_distribution(10000, 0, 1000, "0.01", approved=True) == 0
    assert capped_distribution(10000, 10000, 1000, "0.01", approved=True, active=False) == 0
    assert capped_distribution(10000, 10000, 1000, "0.01", approved=True, contract_active=False) == 0


def test_solvency_monitor_uses_thresholds_and_requires_obligations():
    ratio, health, _ = solvency_status(15_000_000, 10_000_000, 0, 0)
    assert ratio == 1.5
    assert health == "HEALTHY"
    assert solvency_status(13_000_000, 10_000_000, 0, 0)[1] == "WATCH"
    assert solvency_status(11_000_000, 10_000_000, 0, 0)[1] == "CRITICAL"
    assert solvency_status(0, 0, 0, 0)[1] == "INPUT REQUIRED"


def test_fee_preview_is_calculated_not_a_collection():
    fee, net = fee_preview(1000, "0.25")
    assert fee == 250
    assert net == 750


def test_export_has_all_13_sheets_and_marks_scenarios():
    investors = [{
        "investor_id": "INV-TEST", "name_ref": "test-only", "package": "Starter",
        "investment": "300", "join_date": "2026-10-22", "contract_end": "2029-10-21",
        "status": "Active", "staking_opt_in": 0, "created_at": "", "updated_at": "",
    }]
    workbook_bytes = _wb_bytes(
        DEFAULTS.copy(), investors, [], [], 
        {"currency": "USDT", "assets": "0", "customer_balances": "0",
         "pending_rewards": "0", "pending_withdrawals": "0",
         "health": "INPUT REQUIRED", "action": "Enter verified inputs."},
        {key: value for key, value in DEFAULTS.items() if key.startswith("stress_")},
    )
    wb = load_workbook(BytesIO(workbook_bytes), data_only=False)
    assert wb.sheetnames == [
        "01_Settings", "02_Packages", "03_Investors", "04_AI_Grid_Revenue",
        "05_Treasury", "06_Daily_Distributions", "07_Staking", "08_P2P_Fees",
        "09_Withdrawal_Fees", "10_Ledger", "11_Solvency_Monitor",
        "12_Stress_Test", "13_Dashboard",
    ]
    assert wb["06_Daily_Distributions"]["H4"].data_type == "f"
    assert wb["12_Stress_Test"]["B17"].value == "=B15-B16"
    assert "not live" in wb["01_Settings"]["A2"].value.lower() or \
           "planning" in wb["01_Settings"]["A2"].value.lower()
