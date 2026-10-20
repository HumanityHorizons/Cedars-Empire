"""Persistent action/evidence grids and calendar reminders for Humanity Horizons."""
from __future__ import annotations

import io
import json
import re
import textwrap
from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo

import pandas as pd
import pymupdf
import streamlit as st
from docx import Document
from docx.shared import Inches, Pt

GRID_COLUMNS = ["Action / requirement", "Status", "Evidence reference", "Owner", "Due date", "Notes"]
STATUSES = ["Not started", "In progress", "Met", "Not applicable", "Blocked"]

PAGE_ACTIONS = {
    "Dashboard": ["Review priority tasks", "Confirm due dates and status", "Record next action"],
    "My Tasks": ["Create or update task", "Assign status, priority and due date", "Confirm completion evidence"],
    "Opportunity Analysis": ["Applicant eligibility", "Deadline and time zone", "Geographic and applicant restrictions", "Funding amount and eligible costs", "Mandatory attachments", "Evaluation criteria and scoring", "Submission channel and format", "Safeguarding, inclusion and data protection", "Reporting and implementation obligations"],
    "Proposal Studio": ["Mirror mandatory donor headings", "Map problem to objectives and results", "Link indicators to evidence and budget", "Flag unsupported claims and missing inputs", "Review safeguarding and sustainability"],
    "CV / Resume Studio": ["Verify dates, credentials and experience", "Check target role requirements", "Review final document formatting"],
    "SDG Navigator": ["Identify SDG and target", "Link proposed outcome to evidence", "Define measurable indicator"],
    "Noticeboard": ["Verify official source URL", "Check deadline and eligibility", "Flag listing for moderation if needed"],
    "HH Standards & Impact": ["Link result to indicator", "Record baseline, target and source", "Review risk and data-quality controls"],
    "Donor Feedback": ["Classify feedback", "Record response owner and status", "Document resolution evidence"],
    "B2B Collaboration": ["Confirm purpose and counterpart", "Record agreed deliverable", "Track next action"],
    "HR & Volunteers": ["Verify role and contact permissions", "Record availability and status", "Track follow-up"],
    "Revenue & Sustainability": ["Label idea versus verified revenue", "Record buyer/contract evidence", "Track costs and next validation step"],
    "Report Automation": ["Validate indicator definitions", "Check reporting period and source data", "Confirm schedule and recipient"],
    "Privacy and Data": ["Review data access and consent", "Check export/download purpose", "Record privacy follow-up"],
    "Admin Console": ["Review audit events", "Check unresolved feedback and tasks", "Export only authorized records"],
    "Calendar & Events": ["Review upcoming events", "Confirm alarm and timezone", "Close completed follow-ups"],
}

def base_rows(page: str) -> list[dict]:
    return [{"Action / requirement": a, "Status": "Not started", "Evidence reference": "", "Owner": "", "Due date": "", "Notes": ""}
            for a in PAGE_ACTIONS.get(page, ["Check page inputs", "Run the primary action", "Verify output and evidence"])]

def _docx_bytes(title: str, frame: pd.DataFrame) -> bytes:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.75); sec.bottom_margin = Inches(0.75)
    sec.left_margin = Inches(0.8); sec.right_margin = Inches(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"; normal.font.size = Pt(10)
    normal.paragraph_format.line_spacing = 1.5
    title_p = doc.add_heading(title, 0); title_p.paragraph_format.line_spacing = 1.5
    p = doc.add_paragraph(f"Generated: {date.today().isoformat()}")
    p.paragraph_format.line_spacing = 1.5
    table = doc.add_table(rows=1, cols=len(frame.columns)); table.style = "Table Grid"
    for i, col in enumerate(frame.columns): table.rows[0].cells[i].text = str(col)
    for _, row in frame.iterrows():
        cells = table.add_row().cells
        for i, value in enumerate(row.tolist()): cells[i].text = "" if pd.isna(value) else str(value)
    out = io.BytesIO(); doc.save(out); return out.getvalue()

def _pdf_bytes(title: str, frame: pd.DataFrame) -> bytes:
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    margin, y = 45, 54
    page.insert_text((margin, y), title[:100], fontsize=16, fontname="helv"); y += 28
    page.insert_text((margin, y), f"Humanity Horizons | Generated {date.today().isoformat()}", fontsize=8.5, fontname="helv"); y += 26
    for idx, row in frame.iterrows():
        line = f"{idx + 1}. " + " | ".join(f"{col}: {'' if pd.isna(row[col]) else str(row[col])}" for col in frame.columns)
        for part in textwrap.wrap(line, width=90, break_long_words=True) or [""]:
            if y > 790:
                page = doc.new_page(width=595, height=842); y = 54
            page.insert_text((margin, y), part, fontsize=8.5, fontname="helv"); y += 13
        y += 5
    count = len(doc)
    for i in range(count):
        doc[i].insert_text((45, 823), "Humanity Horizons | People, Planet and Peace", fontsize=8, fontname="helv")
        doc[i].insert_text((495, 823), f"Page {i+1} of {count}", fontsize=8, fontname="helv")
    result = doc.tobytes(); doc.close(); return result

def render_action_grid(page: str, user_email: str, db_all, db_exec, record_event) -> None:
    key = "hh_live_grid_" + re.sub(r"[^a-z0-9]+", "_", page.lower()).strip("_")
    with st.expander("Action & evidence grid", expanded=(page == "Opportunity Analysis")):
        stored = db_all("SELECT grid_json FROM workspace_grids WHERE owner_email=? AND page_name=?", (user_email, page))
        if key not in st.session_state:
            try:
                rows = json.loads(stored[0]["grid_json"]) if stored and stored[0].get("grid_json") else base_rows(page)
            except Exception:
                rows = base_rows(page)
            st.session_state[key] = pd.DataFrame(rows, columns=GRID_COLUMNS)
        editor_key = key + "_editor"
        if page == "Opportunity Analysis":
            source_text = str(st.session_state.get("source_text", "") or "")
            if st.button("Find source evidence candidates", key=key + "_find_evidence", disabled=not source_text):
                patterns = {
                    "Applicant eligibility": r"eligib(?:le|ility)|who can apply|applicant(?:s)? must|registered organization|eligible applicant",
                    "Deadline and time zone": r"deadline|closing date|due date|submit.{0,35}by|time zone|timezone",
                    "Geographic and applicant restrictions": r"eligible countr(?:y|ies)|geographic(?:al)? focus|country|countries|region",
                    "Funding amount and eligible costs": r"funding available|grant size|maximum award|budget limit|eligible costs|ineligible costs|amount of funding",
                    "Mandatory attachments": r"required documents|supporting documents|attachments?|letter of support|CVs?",
                    "Evaluation criteria and scoring": r"evaluation criteria|selection criteria|scoring|assessment criteria|weighted",
                    "Submission channel and format": r"submission portal|submit.{0,25}application|email.{0,30}application|online application",
                    "Safeguarding, inclusion and data protection": r"safeguarding|child protection|data protection|inclusion|gender|disability",
                    "Reporting and implementation obligations": r"reporting requirements|progress report|final report|implementation period|project duration|monitoring and evaluation",
                }
                source_name = str(st.session_state.get("source_name", "uploaded source"))
                updated = st.session_state[key].copy()
                for idx, row in updated.iterrows():
                    label = str(row.get("Action / requirement", ""))
                    pattern = next((pat for action, pat in patterns.items() if action.lower() in label.lower()), None)
                    match = re.search(pattern, source_text, flags=re.I | re.S) if pattern else None
                    if match:
                        start_pos, end_pos = match.span()
                        lo, hi = max(0, start_pos - 100), min(len(source_text), end_pos + 180)
                        snippet = " ".join(source_text[lo:hi].split())
                        updated.at[idx, "Evidence reference"] = f"{source_name} · character offsets {lo}–{hi}"
                        updated.at[idx, "Notes"] = f"Candidate only — verify against official source: {snippet[:350]}"
                        if str(updated.at[idx, "Status"]) == "Not started":
                            updated.at[idx, "Status"] = "In progress"
                    else:
                        updated.at[idx, "Notes"] = "No matching evidence phrase found automatically; inspect the full official source."
                st.session_state[key] = updated
                st.session_state.pop(editor_key, None)
                st.rerun()
        edited = st.data_editor(
            st.session_state[key], key=editor_key, num_rows="dynamic", use_container_width=True, hide_index=True,
            column_config={
                "Action / requirement": st.column_config.TextColumn(required=True),
                "Status": st.column_config.SelectboxColumn(options=STATUSES, required=True),
                "Evidence reference": st.column_config.TextColumn(help="Official URL, page/section, document name or record ID"),
                "Owner": st.column_config.TextColumn(),
                "Due date": st.column_config.TextColumn(help="YYYY-MM-DD"),
                "Notes": st.column_config.TextColumn(),
            })
        st.session_state[key] = edited.copy()
        total = len(edited); met = int((edited["Status"] == "Met").sum()) if total else 0
        c1, c2, c3 = st.columns(3)
        c1.metric("Actions", total); c2.metric("Met", met); c3.metric("Open", total-met)
        if st.button("Save grid", key=key+"_save", type="primary"):
            payload = edited.fillna("").to_dict(orient="records")
            if stored:
                db_exec("UPDATE workspace_grids SET grid_json=?,updated_at=? WHERE owner_email=? AND page_name=?",
                        (json.dumps(payload, ensure_ascii=False), datetime.now().isoformat(timespec="seconds"), user_email, page))
            else:
                db_exec("INSERT INTO workspace_grids (id,owner_email,page_name,grid_json,created_at,updated_at) VALUES (?,?,?,?,?,?)",
                        (f"{user_email}:{page}", user_email, page, json.dumps(payload, ensure_ascii=False), datetime.now().isoformat(timespec="seconds"), datetime.now().isoformat(timespec="seconds")))
            record_event(user_email, f"action_grid_saved:{page}")
            st.success("Grid saved.")
            st.rerun()
        if total and ((edited["Status"] == "Met") & edited["Evidence reference"].fillna("").astype(str).str.strip().eq("")).any():
            st.warning("A row marked Met has no evidence reference. Verify before treating it as complete.")
        stem = key
        if st.download_button("Export grid CSV", edited.to_csv(index=False).encode("utf-8"), f"{stem}.csv", "text/csv", key=key+"_csv"):
            record_event(user_email, f"action_grid_download_csv:{page}")
        if st.download_button("Export grid Word", _docx_bytes(f"Humanity Horizons — {page} Action Grid", edited), f"{stem}.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", key=key+"_docx"):
            record_event(user_email, f"action_grid_download_docx:{page}")
        if st.download_button("Export grid PDF", _pdf_bytes(f"Humanity Horizons — {page} Action Grid", edited), f"{stem}.pdf", "application/pdf", key=key+"_pdf"):
            record_event(user_email, f"action_grid_download_pdf:{page}")

def render_calendar(user_email: str, is_admin: bool, db_all, db_exec, record_event, full_page: bool = False) -> None:
    """Persistent event scheduler with in-app due/overdue alerts; external delivery is separate."""
    zone_options = ["Africa/Kampala", "UTC", "Africa/Nairobi", "Europe/London"]
    scope = st.container() if full_page else st.expander("Calendar, event schedule & alarms", expanded=False)
    with scope:
        if full_page:
            st.header("Calendar & Events")
        now_utc = datetime.now(ZoneInfo("UTC"))
        local_zone = ZoneInfo("Africa/Kampala")
        now_local = now_utc.astimezone(local_zone)
        events = db_all(
            "SELECT id,title,event_date,event_time,timezone_name,reminder_minutes,status,details,owner_email FROM calendar_events WHERE owner_email=? ORDER BY event_date,event_time",
            (user_email,))
        active = [e for e in events if e.get("status") != "Completed"]
        due_soon = []
        for event in active:
            try:
                tz = ZoneInfo(event.get("timezone_name") or "Africa/Kampala")
                dt = datetime.fromisoformat(f"{event['event_date']}T{event.get('event_time') or '09:00'}").replace(tzinfo=tz)
                minutes = int(event.get("reminder_minutes") or 0)
                delta = (dt - now_utc.astimezone(tz)).total_seconds() / 60
                if -60 <= delta <= minutes:
                    due_soon.append((event, delta))
            except Exception:
                continue
        if due_soon:
            for event, delta in due_soon[:5]:
                when = "due/overdue" if delta <= 0 else f"in about {int(delta)} min"
                st.warning(f"⏰ {event['title']} — {when}. {event['event_date']} {event.get('event_time') or ''} ({event.get('timezone_name') or 'Africa/Kampala'}).")
        if full_page:
            st.caption(f"Local time: {now_local.strftime('%Y-%m-%d %H:%M')} · Africa/Kampala")
        with st.form(("calendar_event_form_full" if full_page else "calendar_event_form_shared") + "_" + re.sub(r"[^a-z0-9]+", "_", user_email.lower())):
            c1, c2 = st.columns(2)
            title = c1.text_input("Event / task title")
            event_date = c2.date_input("Date", value=date.today())
            c3, c4 = st.columns(2)
            event_time = c3.time_input("Time", value=datetime.now().replace(hour=9, minute=0, second=0, microsecond=0).time())
            timezone_name = c4.selectbox("Time zone", zone_options, index=0)
            c5, c6 = st.columns(2)
            reminder = c5.selectbox("Alarm reminder", [("At event", 0), ("5 minutes before", 5), ("15 minutes before", 15), ("30 minutes before", 30), ("1 hour before", 60), ("1 day before", 1440)], format_func=lambda x: x[0], index=2)
            status = c6.selectbox("Status", ["Scheduled", "In progress", "Completed"])
            details = st.text_input("Notes / meeting link")
            save = st.form_submit_button("Save event", type="primary")
        if save:
            if not title.strip():
                st.error("Enter an event title.")
            else:
                eid = str(__import__("uuid").uuid4())
                db_exec("INSERT INTO calendar_events (id,owner_email,title,event_date,event_time,timezone_name,reminder_minutes,status,details,created_at,updated_at) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                        (eid, user_email, title.strip(), event_date.isoformat(), event_time.strftime("%H:%M"), timezone_name, int(reminder[1]), status, details.strip(), datetime.now().isoformat(timespec="seconds"), datetime.now().isoformat(timespec="seconds")))
                record_event(user_email, "calendar_event_created")
                st.success("Event scheduled.")
                st.rerun()
        st.subheader("Upcoming events")
        visible = [e for e in events if e.get("status") != "Completed"]
        if visible:
            frame = pd.DataFrame([{"Event": e["title"], "Date": e["event_date"], "Time": e.get("event_time"), "Time zone": e.get("timezone_name"), "Reminder (min)": e.get("reminder_minutes"), "Status": e.get("status"), "Notes / link": e.get("details")} for e in visible])
            st.dataframe(frame, use_container_width=True, hide_index=True)
            options = {f"{e['event_date']} {e.get('event_time') or ''} · {e['title']}": e["id"] for e in visible}
            chosen = st.selectbox("Update event", list(options), key=("calendar_event_update_full" if full_page else "calendar_event_update_shared") + "_" + re.sub(r"[^a-z0-9]+", "_", user_email.lower()))
            c1, c2 = st.columns(2)
            if c1.button("Mark completed", key=("calendar_done_full" if full_page else "calendar_done_shared") + "_" + re.sub(r"[^a-z0-9]+", "_", user_email.lower())):
                db_exec("UPDATE calendar_events SET status=?,updated_at=? WHERE id=? AND owner_email=?", ("Completed", datetime.now().isoformat(timespec="seconds"), options[chosen], user_email))
                record_event(user_email, "calendar_event_completed"); st.rerun()
            if c2.button("Delete event", key=("calendar_delete_full" if full_page else "calendar_delete_shared") + "_" + re.sub(r"[^a-z0-9]+", "_", user_email.lower())):
                db_exec("DELETE FROM calendar_events WHERE id=? AND owner_email=?", (options[chosen], user_email))
                record_event(user_email, "calendar_event_deleted"); st.rerun()
        else:
            st.info("No upcoming events.")
        st.caption("In-app alarms appear while this app is open and refreshing. Alerts while the app is closed need a configured scheduled email or push-notification service.")

def render_document_exports(title: str, text: str, key: str, record_event, user_email: str) -> None:
    if not text or not text.strip():
        return
    stem = re.sub(r"[^a-z0-9]+", "_", title.lower()).strip("_")[:70] or "humanity_horizons_document"
    doc = Document()
    sec = doc.sections[0]; sec.top_margin=Inches(0.8); sec.bottom_margin=Inches(0.8); sec.left_margin=Inches(0.9); sec.right_margin=Inches(0.9)
    normal=doc.styles["Normal"]; normal.font.name="Arial"; normal.font.size=Pt(10.5); normal.paragraph_format.line_spacing=1.5
    title_p=doc.add_heading(title,0); title_p.paragraph_format.line_spacing=1.5
    for raw in text.splitlines():
        line=raw.strip()
        if not line: doc.add_paragraph(""); continue
        if len(line)<90 and (line.isupper() or line.endswith(":")):
            p=doc.add_heading(line.rstrip(":"),level=1); p.paragraph_format.line_spacing=1.5
        else:
            p=doc.add_paragraph(line); p.paragraph_format.line_spacing=1.5
    word=io.BytesIO(); doc.save(word)
    pdf=pymupdf.open(); pg=pdf.new_page(width=595,height=842); margin,y=54,58
    pg.insert_text((margin,y),title[:100],fontsize=17,fontname="helv"); y+=30
    pg.insert_text((margin,y),f"Humanity Horizons | Generated {date.today().isoformat()}",fontsize=8.5,fontname="helv"); y+=28
    for raw in text.splitlines():
        line=raw.strip()
        if not line: y+=12; continue
        heading=len(line)<90 and (line.isupper() or line.endswith(":"))
        for part in textwrap.wrap(line,width=75 if heading else 90,break_long_words=True) or [""]:
            if y>790: pg=pdf.new_page(width=595,height=842); y=58
            pg.insert_text((margin,y),part,fontsize=11 if heading else 9.5,fontname="helv"); y+=(11 if heading else 9.5)*1.5
        y+=5
    n=len(pdf)
    for i in range(n):
        pdf[i].insert_text((54,822),"Humanity Horizons | People, Planet and Peace",fontsize=8,fontname="helv")
        pdf[i].insert_text((500,822),f"Page {i+1} of {n}",fontsize=8,fontname="helv")
    pdfbytes=pdf.tobytes(); pdf.close()
    if st.download_button("Download Word (.docx)",word.getvalue(),f"{stem}.docx","application/vnd.openxmlformats-officedocument.wordprocessingml.document",key=key+"_word"):
        record_event(user_email,f"document_download_docx:{title}")
    if st.download_button("Download PDF",pdfbytes,f"{stem}.pdf","application/pdf",key=key+"_pdf"):
        record_event(user_email,f"document_download_pdf:{title}")
