"""Admin-only Cedars Empire treasury model. Planning and reconciliation tools only; no payment execution."""
from __future__ import annotations

from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO
import uuid

import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

PREFIX = "cedars_treasury_model:"
PACKAGES = {"Starter": Decimal("300"), "Bronze": Decimal("1000"), "Silver": Decimal("2500"),
            "Gold": Decimal("5000"), "Platinum": Decimal("10000"), "Diamond": Decimal("25000")}
GRIDS = ["NeuroWave", "EchoMap", "BioLoop", "DreamGrid", "EchoPulse", "SolarGrid", "HydroGrid",
         "Cleaning", "Shelter & Accommodation", "Energy & Utilities", "AI Training & Data"]
DEFAULTS = {
    "contract_days": "1095", "daily_distribution": "0.01", "customer_staking": "0.15",
    "treasury_compounding": "0.05", "treasury_staking": "0.06", "reserve_lock": "0.40",
    "withdrawal_fee": "0.25", "p2p_fee": "0.25", "solvency_threshold": "1.20",
    "healthy_threshold": "1.50", "distribution_approval": "Not approved",
    "distribution_pool": "0",
    "solvency_assets": "0", "solvency_customer_balances": "0",
    "solvency_pending_rewards": "0", "solvency_pending_withdrawals": "0",
    "solvency_currency": "USDT", "stress_opening": "0", "stress_grid_revenue": "0",
    "stress_p2p_fees": "0", "stress_withdrawal_fees": "0", "stress_compounding": "0",
    "stress_staking": "0", "stress_other_income": "0", "stress_distributions": "0",
    "stress_withdrawals": "0", "stress_operations": "0",
}
CURRENCIES = ["USDT", "USD", "UGX", "TRX", "BNB", "EUR", "Other"]
MODEL_WARNING = (
    "TEST / PLANNING MODE ONLY. No wallet signing, deposit acceptance, staking provider, scheduled payout, "
    "bank/chain feed, or automated withdrawal is connected. These calculations do not establish actual assets, "
    "customer liabilities, legal compliance, solvency, or an investment return."
)

def dec(value, default="0") -> Decimal:
    try:
        return Decimal(str(value if value not in (None, "") else default))
    except (InvalidOperation, ValueError, TypeError):
        return Decimal(default)

def capped_distribution(investment, total_active_capital, available_pool, daily_rate, approved=False,
                        active=True, contract_active=True) -> Decimal:
    """Return a model-only capped payout. This function never sends or books money."""
    inv, capital, pool, rate = map(dec, (investment, total_active_capital, available_pool, daily_rate))
    if not (approved and active and contract_active and inv > 0 and capital > 0 and pool > 0 and rate >= 0):
        return Decimal("0")
    return min(inv * rate, pool * (inv / capital))

def solvency_status(assets, customer_balances, pending_rewards, pending_withdrawals,
                    watch_threshold="1.20", healthy_threshold="1.50") -> tuple[Decimal, str, str]:
    obligations = sum((dec(x) for x in (customer_balances, pending_rewards, pending_withdrawals)), Decimal("0"))
    ratio = dec(assets) / obligations if obligations > 0 else Decimal("0")
    if obligations <= 0:
        return ratio, "INPUT REQUIRED", "Enter independently reconciled assets and obligations."
    if ratio < dec(watch_threshold):
        return ratio, "CRITICAL", "Recommend pausing/reducing distributions pending authorized human review."
    if ratio < dec(healthy_threshold):
        return ratio, "WATCH", "Enhanced reconciliation and human review recommended."
    return ratio, "HEALTHY", "Continue independent reconciliation; no automatic action is enabled."

def fee_preview(amount, rate) -> tuple[Decimal, Decimal]:
    gross, fee_rate = dec(amount), dec(rate)
    if gross < 0 or fee_rate < 0 or fee_rate > 1:
        raise ValueError("Amount must be non-negative and fee rate must be between 0% and 100%.")
    fee = gross * fee_rate
    return fee, gross - fee

def ensure_tables(db_exec) -> None:
    db_exec("""CREATE TABLE IF NOT EXISTS cedars_model_investors (
        investor_id TEXT PRIMARY KEY, name_ref TEXT NOT NULL, package TEXT NOT NULL,
        investment TEXT NOT NULL, join_date TEXT NOT NULL, contract_end TEXT NOT NULL,
        status TEXT NOT NULL, staking_opt_in INTEGER NOT NULL DEFAULT 0,
        created_at TEXT NOT NULL, updated_at TEXT NOT NULL)""")
    db_exec("""CREATE TABLE IF NOT EXISTS cedars_model_grid_revenue (
        revenue_id TEXT PRIMARY KEY, revenue_date TEXT NOT NULL, grid_name TEXT NOT NULL,
        amount TEXT NOT NULL, currency TEXT NOT NULL, status TEXT NOT NULL,
        evidence_ref TEXT, notes TEXT, created_by TEXT NOT NULL,
        created_at TEXT NOT NULL, updated_at TEXT NOT NULL)""")

def load_settings(db_all) -> dict:
    values = dict(DEFAULTS)
    for row in db_all("SELECT setting_key,setting_value FROM app_settings"):
        key = str(row.get("setting_key", ""))
        if key.startswith(PREFIX):
            values[key[len(PREFIX):]] = str(row.get("setting_value", ""))
    return values

def save_settings(db_exec, record_event, user_email: str, values: dict) -> None:
    now = datetime.utcnow().replace(microsecond=0).isoformat() + "Z"
    for key, value in values.items():
        db_exec("""INSERT INTO app_settings (setting_key,setting_value,updated_by,updated_at)
                   VALUES (?,?,?,?) ON CONFLICT (setting_key) DO UPDATE
                   SET setting_value=excluded.setting_value,updated_by=excluded.updated_by,updated_at=excluded.updated_at""",
                (PREFIX + key, str(value), user_email, now))
    record_event(user_email, "treasury_model_settings_updated", "planning assumptions only")

def _wb_bytes(settings: dict, investors: list[dict], revenue: list[dict], ledger: list[dict],
              solvency: dict, stress: dict) -> bytes:
    """Create a portable 13-sheet planning workbook snapshot; this is not a transaction file."""
    wb = Workbook()
    sheet_names = ["01_Settings","02_Packages","03_Investors","04_AI_Grid_Revenue","05_Treasury",
                   "06_Daily_Distributions","07_Staking","08_P2P_Fees","09_Withdrawal_Fees",
                   "10_Ledger","11_Solvency_Monitor","12_Stress_Test","13_Dashboard"]
    ws = wb.active
    ws.title = sheet_names[0]
    for name in sheet_names[1:]:
        wb.create_sheet(name)
    navy, blue, white, yellow = "0A2540", "0066CC", "FFFFFF", "FFF2CC"
    def format_sheet(sheet, title, note, headers):
        n = len(headers)
        sheet.append([title])
        sheet.merge_cells(start_row=1,start_column=1,end_row=1,end_column=n)
        sheet.cell(1,1).fill = PatternFill("solid", fgColor=navy)
        sheet.cell(1,1).font = Font(name="Arial",size=14,bold=True,color=white)
        sheet.append([note])
        sheet.merge_cells(start_row=2,start_column=1,end_row=2,end_column=n)
        sheet.cell(2,1).alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[2].height = 30
        sheet.append(headers)
        for c in range(1,n+1):
            cell = sheet.cell(3,c)
            cell.fill = PatternFill("solid",fgColor=blue)
            cell.font = Font(bold=True,color=white)
            cell.alignment = Alignment(wrap_text=True,vertical="center")
        sheet.row_dimensions[3].height = 32
        sheet.freeze_panes = "A4"
        sheet.sheet_view.showGridLines = False
        for c in range(1,n+1):
            sheet.column_dimensions[get_column_letter(c)].width = min(36,max(14,len(str(headers[c-1]))+3))
    def append_rows(sheet, rows):
        for row in rows:
            sheet.append([row.get(k,"") for k in sheet.headers] if hasattr(sheet,"headers") else row)
    note = "Planning/test export only. Not live wallet data, a guarantee, a solvency certification, or a payment instruction."
    # Settings
    s = wb["01_Settings"]; format_sheet(s,"CEDARS EMPIRE TREASURY MODEL v1.0",note,["Setting","Value","Unit / caution"])
    settings_rows = [
        ("Contract Days", settings["contract_days"], "days"),
        ("Daily Distribution Target", settings["daily_distribution"], "rate/day — illustrative, not promised"),
        ("Customer Annual Staking", settings["customer_staking"], "rate/year — illustrative"),
        ("Treasury Monthly Compounding", settings["treasury_compounding"], "rate/month — not actual yield"),
        ("Treasury Monthly Staking", settings["treasury_staking"], "rate/month — no provider connected"),
        ("Reserve Lock", settings["reserve_lock"], "share — cannot lock on-chain funds"),
        ("Withdrawal Fee", settings["withdrawal_fee"], "proposed rate; review terms and law"),
        ("P2P Fee", settings["p2p_fee"], "proposed rate; review terms and law"),
        ("Solvency Threshold", settings["solvency_threshold"], "coverage multiple"),
        ("Healthy Threshold", settings["healthy_threshold"], "coverage multiple"),
        ("Model Approval", settings["distribution_approval"], "planning only; not payment authorization"),
        ("Model As-Of Date", date.today().isoformat(), "local workbook creation date"),
        ("Distribution Pool (Preview Only)", settings.get("distribution_pool", "0"), "Planning only; not authorized to pay."),
    ]
    for row in settings_rows: s.append(list(row))
    for r in range(4,17): s.cell(r,2).fill=PatternFill("solid",fgColor=yellow); s.cell(r,2).font=Font(color="0000FF")
    # Packages
    s=wb["02_Packages"]; format_sheet(s,"PACKAGE REFERENCE",note,["Package","Investment reference amount","Caution"])
    for name, amount in PACKAGES.items(): s.append([name,float(amount),"Test package reference only; not a receipt or investment offer."])
    # Investors
    s=wb["03_Investors"]; format_sheet(s,"INVESTOR REGISTER (MODEL)",note,["Investor ID","Name / Reference","Package","Investment","Join Date","Contract End","Remaining Days","Status","Staking Opt-In"])
    for r, inv in enumerate(investors,4):
        s.append([inv.get("investor_id"),inv.get("name_ref"),inv.get("package"),float(dec(inv.get("investment"))),
                  inv.get("join_date"),inv.get("contract_end"),None,inv.get("status"),"Yes" if str(inv.get("staking_opt_in")) in ("1","True","true") else "No"])
        s.cell(r,7,f'=IF(F{r}="","",F{r}-TODAY())')
    # Grid revenue
    s=wb["04_AI_Grid_Revenue"]; format_sheet(s,"AI GRID REVENUE EVIDENCE REGISTER",note,["Date","Grid","Amount","Currency","Evidence Status","Evidence Reference","Notes","Counted as Cash?"])
    for r, item in enumerate(revenue,4):
        s.append([item.get("revenue_date"),item.get("grid_name"),float(dec(item.get("amount"))),item.get("currency"),item.get("status"),item.get("evidence_ref"),item.get("notes"),"NO — match to Treasury Ledger"])
    # Treasury roll-forward — separate actual ledger balances from scenario yields
    s=wb["05_Treasury"]; format_sheet(s,"TREASURY LEDGER SNAPSHOT",note,["Currency","Admin-Entered Ledger Balance","Modeled Compounding (not actual)","Modeled Staking (not actual)","Reserve Target (modeled)","Available Pool (modeled)","Note"])
    by_currency={}
    for item in ledger:
        c=str(item.get("currency","Other"))
        amount=dec(item.get("amount"))
        typ=str(item.get("entry_type",""))
        if typ=="Verified income received": by_currency[c]=by_currency.get(c,Decimal(0))+amount
        elif typ in ("Verified payment sent","Verified expense","Refund"): by_currency[c]=by_currency.get(c,Decimal(0))-amount
    for c, bal in sorted(by_currency.items()):
        s.append([c,float(bal),None,None,None,None,"Admin-entered ledger net only; not independently live-verified."])
        r=s.max_row
        s.cell(r,3,f"=B{r}*'01_Settings'!$B$7")
        s.cell(r,4,f"=B{r}*'01_Settings'!$B$8")
        s.cell(r,5,f"=B{r}*'01_Settings'!$B$9")
        s.cell(r,6,f'=MAX(0,B{r}-E{r})')
    # Distribution calculator
    s=wb["06_Daily_Distributions"]; format_sheet(s,"DAILY DISTRIBUTION PREVIEW",note,["Investor ID","Package","Investment","Status","Weight","Gross Target","Pool Share","Capped Distribution Preview"])
    active=[x for x in investors if x.get("status")=="Active"]
    cap=sum((dec(x.get("investment")) for x in active),Decimal(0))
    for r, inv in enumerate(investors,4):
        investment=dec(inv.get("investment"))
        s.append([inv.get("investor_id"),inv.get("package"),float(investment),inv.get("status"),None,None,None,None])
        s.cell(r,5,f'=IF(AND(D{r}="Active",SUMIF($D$4:$D$503,"Active",$C$4:$C$503)>0),C{r}/SUMIF($D$4:$D$503,"Active",$C$4:$C$503),0)')
        s.cell(r,6,f'=IF(D{r}="Active",C{r}*\'01_Settings\'!$B$5,0)')
        s.cell(r,7,f'=IF(D{r}="Active",\'01_Settings\'!$B$16*E{r},0)')
        s.cell(r,8,f'=IF(\'01_Settings\'!$B$14="Approved",MIN(F{r},G{r}),0)')
    # Customer staking (illustrative, no accrual booked)
    s=wb["07_Staking"]; format_sheet(s,"CUSTOMER STAKING ACCRUAL PREVIEW",note,["Investor ID","Investment","Status","Opt-In","Annual Rate","Daily Accrual","Monthly Accrual","Caution"])
    for r, inv in enumerate(investors,4):
        investment=dec(inv.get("investment")); opted=str(inv.get("staking_opt_in")) in ("1","True","true")
        s.append([inv.get("investor_id"),float(investment),inv.get("status"),"Yes" if opted else "No",None,None,None,"Illustrative only; not a booked liability or payment."])
        s.cell(r,5,"='01_Settings'!$B$6")
        s.cell(r,6,f'=IF(AND(C{r}="Active",D{r}="Yes"),B{r}*E{r}/365,0)')
        s.cell(r,7,f'=IF(AND(C{r}="Active",D{r}="Yes"),B{r}*E{r}/12,0)')
    # Fee calculators
    s=wb["08_P2P_Fees"]; format_sheet(s,"P2P FEE CALCULATOR",note,["Date","Transfer Ref","Sender Ref","Receiver Ref","Amount","Fee Rate","Fee","Net Received","Status","Evidence"])
    for r in range(4,24): s.cell(r,6,"='01_Settings'!$B$11"); s.cell(r,7,f'=E{r}*F{r}'); s.cell(r,8,f'=E{r}-G{r}')
    s=wb["09_Withdrawal_Fees"]; format_sheet(s,"WITHDRAWAL FEE CALCULATOR",note,["Date","Withdrawal Ref","Investor ID","Amount","Fee Rate","Fee","Net Paid","Status","Evidence"])
    for r in range(4,24): s.cell(r,5,"='01_Settings'!$B$10"); s.cell(r,6,f'=D{r}*E{r}'); s.cell(r,7,f'=D{r}-F{r}')
    # Detailed ledger
    s=wb["10_Ledger"]; format_sheet(s,"LEDGER EXPORT",note,["Entry Type","Amount","Currency","Source","Reference","Notes","Entered By","Created At"])
    for item in ledger:
        s.append([item.get("entry_type"),float(dec(item.get("amount"))),item.get("currency"),item.get("source"),item.get("reference"),item.get("notes"),item.get("entered_by"),item.get("created_at")])
    # Solvency
    s=wb["11_Solvency_Monitor"]; format_sheet(s,"SOLVENCY MONITOR — MANUAL INPUTS",note,["Currency","Verified Assets Input","Customer Balances","Pending Rewards","Pending Withdrawals","Total Obligations","Coverage Ratio","Health","Action"])
    s.append([solvency.get("currency","USDT"),float(dec(solvency.get("assets"))),float(dec(solvency.get("customer_balances")),),float(dec(solvency.get("pending_rewards"))),float(dec(solvency.get("pending_withdrawals"))),None,None,solvency.get("health","INPUT REQUIRED"),solvency.get("action","Enter verified inputs.")])
    s["F4"]="=SUM(C4:E4)"; s["G4"]='=IF(F4<=0,0,B4/F4)'
    # Stress test
    s=wb["12_Stress_Test"]; format_sheet(s,"24-MONTH ZERO-NEW-DEPOSIT STRESS TEST",note,["Scenario Input","Amount","Interpretation"])
    for key,label in [("stress_opening","Opening Treasury"),("stress_grid_revenue","Grid Revenue"),("stress_p2p_fees","P2P Fees"),("stress_withdrawal_fees","Withdrawal Fees"),("stress_compounding","Scenario Compounding"),("stress_staking","Scenario Staking"),("stress_other_income","Other Income"),("stress_distributions","Distributions"),("stress_withdrawals","Withdrawals"),("stress_operations","Operations")]:
        s.append([label,float(dec(stress.get(key))),"User-entered scenario value; not an actual or guaranteed flow."])
    s.append(["New Deposits",0,"Fixed to zero for the scenario."])
    s.append(["Total Cash-In","=SUM(B5:B10)","Illustrative only; opening balance is not cash-in."])
    s.append(["Total Cash-Out","=SUM(B11:B13)","Illustrative only."])
    s.append(["Net Scenario Result","=B15-B16","Illustrative only."])
    # Dashboard
    s=wb["13_Dashboard"]; format_sheet(s,"CEDARS EMPIRE TREASURY DASHBOARD",note,["Metric","Value","Status / Caveat"])
    dashboard=[
        ["Model-register investors",len(investors),"Not proof of funded accounts."],
        ["Grid revenue evidence entries",len(revenue),"Not verified cash."],
        ["Ledger entries",len(ledger),"Admin-entered; not independent live chain verification."],
        ["Solvency coverage ratio","='11_Solvency_Monitor'!G4",solvency.get("health","INPUT REQUIRED")],
        ["Distribution approval","='01_Settings'!B14","Planning-only preview; no payment action."],
        ["Payment integration","Not connected","No transfer is made by this workbook."],
    ]
    for row in dashboard: s.append(row)
    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines=False
        sheet.freeze_panes = sheet.freeze_panes or "A4"
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type=="f": cell.font=Font(color="000000")
    bio=BytesIO()
    wb.save(bio)
    return bio.getvalue()

def render_treasury_management(db_all, db_exec, record_event, user_email: str, durable_storage: bool) -> None:
    """Render an admin-only, persistent planning and ledger review workspace."""
    ensure_tables(db_exec)
    st.header("Treasury Management Model v1.0")
    st.warning(MODEL_WARNING)
    if durable_storage:
        st.caption("Persistence: managed PostgreSQL is configured for this app. Verify the database backup/restore path during testing.")
    else:
        st.error("Persistence warning: DATABASE_URL is not configured in this app; hosted local SQLite may be lost on redeploy. Do not enter real financial or investor records until persistent storage is verified.")
    settings=load_settings(db_all)
    investor_rows=db_all("SELECT investor_id,name_ref,package,investment,join_date,contract_end,status,staking_opt_in,created_at,updated_at FROM cedars_model_investors ORDER BY created_at DESC")
    revenue_rows=db_all("SELECT revenue_id,revenue_date,grid_name,amount,currency,status,evidence_ref,notes,created_by,created_at FROM cedars_model_grid_revenue ORDER BY revenue_date DESC,created_at DESC")
    ledger_rows=db_all("SELECT entry_type,amount,currency,source,reference,notes,entered_by,created_at FROM treasury_ledger ORDER BY created_at DESC")
    tabs=st.tabs(["01 Settings","02 Packages","03 Investors","04 AI Grid Revenue","05 Treasury","06 Daily Distributions","07 Staking","08 P2P Fees","09 Withdrawal Fees","10 Ledger","11 Solvency Monitor","12 Stress Test","13 Dashboard"])

    with tabs[0]:
        st.subheader("Global assumptions")
        st.caption("Blue/yellow fields are adjustable planning assumptions. Save applies them to this app's model only.")
        with st.form("cedars_model_settings"):
            c1,c2,c3=st.columns(3)
            contract=c1.number_input("Contract days",min_value=1,max_value=3650,value=int(dec(settings["contract_days"],"1095")))
            distribution=c2.number_input("Daily distribution target (%)",min_value=0.0,max_value=100.0,value=float(dec(settings["daily_distribution"])*100),step=0.1)/100
            staking=c3.number_input("Customer annual staking (%)",min_value=0.0,max_value=100.0,value=float(dec(settings["customer_staking"])*100),step=0.5)/100
            c1,c2,c3=st.columns(3)
            compound=c1.number_input("Treasury compounding per month (%)",min_value=0.0,max_value=100.0,value=float(dec(settings["treasury_compounding"])*100),step=0.5)/100
            treasury_stake=c2.number_input("Treasury staking per month (%)",min_value=0.0,max_value=100.0,value=float(dec(settings["treasury_staking"])*100),step=0.5)/100
            reserve=c3.number_input("Reserve target (%)",min_value=0.0,max_value=100.0,value=float(dec(settings["reserve_lock"])*100),step=1.0)/100
            c1,c2,c3=st.columns(3)
            withdrawal_fee=c1.number_input("Withdrawal fee (%)",min_value=0.0,max_value=100.0,value=float(dec(settings["withdrawal_fee"])*100),step=1.0)/100
            p2p_fee=c2.number_input("P2P fee (%)",min_value=0.0,max_value=100.0,value=float(dec(settings["p2p_fee"])*100),step=1.0)/100
            watch=c3.number_input("Critical coverage threshold (x)",min_value=0.0,max_value=10.0,value=float(dec(settings["solvency_threshold"])),step=0.05)
            healthy=st.number_input("Healthy coverage threshold (x)",min_value=0.0,max_value=10.0,value=float(dec(settings["healthy_threshold"])),step=0.05)
            st.caption("Model approvals never enable a transfer; the live payment pathway remains disconnected.")
            save=st.form_submit_button("Save model settings",type="primary")
        if save:
            if healthy < watch:
                st.error("Healthy threshold must be equal to or above the critical coverage threshold.")
            else:
                save_settings(db_exec,record_event,user_email,{
                    "contract_days":contract,"daily_distribution":distribution,"customer_staking":staking,
                    "treasury_compounding":compound,"treasury_staking":treasury_stake,"reserve_lock":reserve,
                    "withdrawal_fee":withdrawal_fee,"p2p_fee":p2p_fee,"solvency_threshold":watch,
                    "healthy_threshold":healthy,
                })
                st.success("Planning settings saved.")
                st.rerun()

    with tabs[1]:
        st.subheader("Fixed reference packages")
        st.dataframe(pd.DataFrame([{"Package":k,"Investment reference":float(v),"Status":"Planning reference only"} for k,v in PACKAGES.items()]),use_container_width=True,hide_index=True)

    with tabs[2]:
        st.subheader("Model investor register")
        st.caption("Use test aliases/references only. A package selection is not evidence that funds were received; this register does not establish customer liabilities.")
        with st.form("cedars_model_add_investor"):
            c1,c2=st.columns(2)
            investor_name=c1.text_input("Investor reference / test alias")
            package=c2.selectbox("Package",list(PACKAGES))
            joined=st.date_input("Join date",value=date.today())
            status=st.selectbox("Model status",["Active","Suspended","Matured","Closed"])
            opted=st.checkbox("Customer staking opt-in recorded (model only)")
            save_investor=st.form_submit_button("Add model investor")
        if save_investor:
            if not investor_name.strip():
                st.error("Enter a test alias/reference; avoid unnecessary personal data.")
            else:
                investor_id="INV-"+uuid.uuid4().hex[:10].upper()
                end=joined+timedelta(days=int(dec(settings["contract_days"],"1095")))
                now=datetime.utcnow().replace(microsecond=0).isoformat()+"Z"
                db_exec("INSERT INTO cedars_model_investors (investor_id,name_ref,package,investment,join_date,contract_end,status,staking_opt_in,created_at,updated_at) VALUES (?,?,?,?,?,?,?,?,?,?)",
                        (investor_id,investor_name.strip(),package,str(PACKAGES[package]),joined.isoformat(),end.isoformat(),status,1 if opted else 0,now,now))
                record_event(user_email,"treasury_model_investor_added","model only: "+investor_id)
                st.success("Model investor saved; no deposit or account funding is implied.")
                st.rerun()
        if investor_rows:
            st.dataframe(pd.DataFrame(investor_rows).drop(columns=["created_at","updated_at"]),use_container_width=True,hide_index=True)
        else: st.info("No model investors registered.")

    with tabs[3]:
        st.subheader("Daily grid-revenue evidence capture")
        with st.form("cedars_model_revenue"):
            c1,c2,c3=st.columns(3)
            revenue_date=c1.date_input("Revenue evidence date",value=date.today())
            grid=c2.selectbox("Grid / revenue stream",GRIDS)
            currency=c3.selectbox("Currency",CURRENCIES)
            amount=st.number_input("Claimed receipt amount",min_value=0.0,value=0.0,step=1.0,format="%.4f")
            evidence_status=st.selectbox("Evidence status",["Unverified","Evidence submitted","Pending review","Reconciled","Rejected"])
            evidence=st.text_input("Invoice / transaction reference / evidence URL")
            notes=st.text_area("Reconciliation notes")
            independent_check=st.checkbox("For Reconciled status: I independently matched the correct network/account, asset, amount and transaction evidence.")
            save_revenue=st.form_submit_button("Save revenue evidence")
        if save_revenue:
            if amount<=0 or not evidence.strip():
                st.error("A positive amount and evidence reference are required.")
            elif evidence_status=="Reconciled" and not independent_check:
                st.error("Tick the independent reconciliation confirmation before using Reconciled status.")
            else:
                now=datetime.utcnow().replace(microsecond=0).isoformat()+"Z"
                db_exec("INSERT INTO cedars_model_grid_revenue (revenue_id,revenue_date,grid_name,amount,currency,status,evidence_ref,notes,created_by,created_at,updated_at) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                        (str(uuid.uuid4()),revenue_date.isoformat(),grid,str(amount),currency,evidence_status,evidence.strip(),notes.strip(),user_email,now,now))
                record_event(user_email,"treasury_model_grid_revenue_recorded",evidence_status)
                st.success("Evidence row saved. It is not added to verified Treasury assets; record/reconcile the actual receipt in the Treasury Ledger separately.")
                st.rerun()
        if revenue_rows: st.dataframe(pd.DataFrame(revenue_rows).drop(columns=["created_by","created_at","updated_at"]),use_container_width=True,hide_index=True)
        else: st.info("No grid revenue evidence rows yet.")

    with tabs[4]:
        st.subheader("Treasury roll-forward by currency")
        st.caption("Balance is computed only from admin-entered ledger entries and never mixes currencies. It is not a live wallet balance.")
        balances={}
        for item in ledger_rows:
            cur=str(item.get("currency") or "Other")
            amt=dec(item.get("amount"))
            kind=str(item.get("entry_type") or "")
            if kind=="Verified income received": balances[cur]=balances.get(cur,Decimal(0))+amt
            elif kind in ("Verified payment sent","Verified expense","Refund"): balances[cur]=balances.get(cur,Decimal(0))-amt
        if balances:
            st.dataframe(pd.DataFrame([{"Currency":c,"Ledger net":float(v),"Data source":"Admin-entered ledger; not independently chain-attested"} for c,v in sorted(balances.items())]),use_container_width=True,hide_index=True)
        else: st.info("No counted ledger movements yet.")
        c1,c2,c3=st.columns(3)
        c1.metric("Grid evidence records",len(revenue_rows))
        c2.metric("Ledger entries",len(ledger_rows))
        c3.metric("Model investors",len(investor_rows))
        st.info("Modeled treasury compounding, treasury staking and reserve targets are shown in the stress test; they are not added to actual ledger assets.")

    with tabs[5]:
        st.subheader("Capped daily distribution preview")
        st.caption("This is a calculation preview only. It does not approve, queue, sign, or send a payout.")
        active=[]
        asof=st.date_input("Distribution preview date",value=date.today(),key="cedars_dist_date")
        for item in investor_rows:
            try:
                active_days=(item["status"]=="Active" and date.fromisoformat(str(item["join_date"])[:10])<=asof and date.fromisoformat(str(item["contract_end"])[:10])>=asof)
            except Exception: active_days=False
            if active_days: active.append(item)
        capital=sum((dec(item["investment"]) for item in active),Decimal(0))
        pool=st.number_input("Manually confirmed available pool for workbook preview (planning units; single currency)",min_value=0.0,value=float(dec(settings.get("distribution_pool","0"))),step=1.0,key="cedars_dist_pool")
        approved=st.checkbox("Save approval flag for workbook preview only (still never authorizes a payment)",value=settings.get("distribution_approval","Not approved")=="Approved",key="cedars_dist_approved")
        if st.button("Save workbook preview inputs",key="cedars_dist_save_preview"):
            save_settings(db_exec,record_event,user_email,{
                "distribution_pool":str(dec(pool)),
                "distribution_approval":"Approved" if approved else "Not approved",
            })
            st.success("Pool and approval flag saved for the downloadable model preview only.")
            st.rerun()
        rows=[]
        for item in active:
            inv=dec(item["investment"])
            payout=capped_distribution(inv,capital,pool,dec(settings["daily_distribution"]),approved=approved)
            rows.append({"Investor ID":item["investor_id"],"Package":item["package"],"Investment reference":float(inv),
                         "Weight":float(inv/capital) if capital else 0.0,"Gross 1-day target (illustrative)":float(inv*dec(settings["daily_distribution"])),
                         "Capped distribution preview":float(payout),"Status":"PREVIEW ONLY — not payable"})
        st.metric("Active model capital",f"{capital:,.2f}")
        st.metric("Total capped distribution preview",f"{sum((dec(x['Capped distribution preview']) for x in rows),Decimal(0)):,.2f}")
        if rows: st.dataframe(pd.DataFrame(rows),use_container_width=True,hide_index=True)
        else: st.info("No active model investors within their contract dates.")

    with tabs[6]:
        st.subheader("Customer staking accrual preview")
        opted=[x for x in investor_rows if x.get("status")=="Active" and str(x.get("staking_opt_in")) in ("1","True","true")]
        staking_rows=[]
        for item in opted:
            investment=dec(item["investment"])
            staking_rows.append({"Investor ID":item["investor_id"],"Investment reference":float(investment),
                "Annual simple accrual (model)":float(investment*dec(settings["customer_staking"])),
                "Monthly simple accrual (model)":float(investment*dec(settings["customer_staking"])/12),
                "Daily simple accrual (model)":float(investment*dec(settings["customer_staking"])/365),
                "Caution":"Not accrued as a liability; not paid."})
        if staking_rows: st.dataframe(pd.DataFrame(staking_rows),use_container_width=True,hide_index=True)
        else: st.info("No active model investors opted into staking.")
        st.caption("No approved staking provider, contract or transaction is connected.")

    with tabs[7]:
        st.subheader("P2P fee calculator")
        p2p_amount=st.number_input("Transfer amount",min_value=0.0,value=1000.0,step=10.0,key="cedars_p2p_amount")
        c1,c2=st.columns(2)
        c1.metric("Illustrative fee",f"{p2p_amount*float(dec(settings['p2p_fee'])):,.2f}")
        c2.metric("Illustrative net received",f"{p2p_amount*(1-float(dec(settings['p2p_fee']))):,.2f}")
        st.caption(f"Configured model fee: {float(dec(settings['p2p_fee'])):.2%}. This is a proposal/calculator; no P2P transfer exists here.")

    with tabs[8]:
        st.subheader("Withdrawal fee calculator")
        withdrawal_amount=st.number_input("Withdrawal amount",min_value=0.0,value=1000.0,step=10.0,key="cedars_withdrawal_amount")
        c1,c2=st.columns(2)
        c1.metric("Illustrative fee",f"{withdrawal_amount*float(dec(settings['withdrawal_fee'])):,.2f}")
        c2.metric("Illustrative net amount",f"{withdrawal_amount*(1-float(dec(settings['withdrawal_fee']))):,.2f}")
        st.caption(f"Configured model fee: {float(dec(settings['withdrawal_fee'])):.2%}. Actual fees require reviewed terms; no withdrawal is sent.")

    with tabs[9]:
        st.subheader("Treasury double-check / reference ledger")
        if ledger_rows: st.dataframe(pd.DataFrame(ledger_rows),use_container_width=True,hide_index=True)
        else: st.info("No ledger entries yet. Add receipts/expenses in the existing Admin Console → Treasury ledger only after evidence reconciliation.")
        st.caption("This model page does not duplicate or post entries to the live ledger.")

    with tabs[10]:
        st.subheader("Solvency monitor — same currency only")
        st.caption("Assets and obligations below must be entered for the same currency and independently reconciled. Do not use modeled investor packages, projected revenue, or assumed yields as assets.")
        with st.form("cedars_solvency_form"):
            cur=st.selectbox("Currency",CURRENCIES,index=CURRENCIES.index(settings.get("solvency_currency","USDT")) if settings.get("solvency_currency","USDT") in CURRENCIES else 0)
            c1,c2=st.columns(2)
            assets=c1.number_input("Verified treasury assets",min_value=0.0,value=float(dec(settings["solvency_assets"])),step=1.0)
            balances_in=c2.number_input("Customer balance obligations",min_value=0.0,value=float(dec(settings["solvency_customer_balances"])),step=1.0)
            c1,c2=st.columns(2)
            pending_rewards=c1.number_input("Pending rewards / liabilities",min_value=0.0,value=float(dec(settings["solvency_pending_rewards"])),step=1.0)
            pending_withdrawals=c2.number_input("Pending withdrawals",min_value=0.0,value=float(dec(settings["solvency_pending_withdrawals"])),step=1.0)
            save_solvency=st.form_submit_button("Save solvency inputs")
        if save_solvency:
            save_settings(db_exec,record_event,user_email,{"solvency_currency":cur,"solvency_assets":assets,
                "solvency_customer_balances":balances_in,"solvency_pending_rewards":pending_rewards,"solvency_pending_withdrawals":pending_withdrawals})
            st.success("Solvency inputs saved for monitoring; this is not a certification.")
            st.rerun()
        ratio,health,action=solvency_status(settings["solvency_assets"],settings["solvency_customer_balances"],
            settings["solvency_pending_rewards"],settings["solvency_pending_withdrawals"],settings["solvency_threshold"],settings["healthy_threshold"])
        st.metric("Coverage ratio",f"{float(ratio):.1%}" if ratio else "Not calculable")
        if health=="HEALTHY": st.success(f"HEALTHY — {action}")
        elif health=="WATCH": st.warning(f"WATCH — {action}")
        elif health=="CRITICAL": st.error(f"CRITICAL — {action}")
        else: st.info(f"INPUT REQUIRED — {action}")
        st.caption("A low ratio does not automatically change or pause any real payout because this page has no payment rail.")

    with tabs[11]:
        st.subheader("24-month zero-new-deposit stress test")
        st.caption("Enter scenario-only assumptions. This is not a forecast unless you replace inputs with evidence-backed, validated operating data.")
        with st.form("cedars_stress_test"):
            defaults={}
            stress_labels=[
                ("stress_opening","Opening treasury"),("stress_grid_revenue","Total grid revenue over 24 months"),
                ("stress_p2p_fees","P2P fees actually collected (scenario)"),("stress_withdrawal_fees","Withdrawal fees actually collected (scenario)"),
                ("stress_compounding","Modeled treasury compounding"),("stress_staking","Modeled treasury staking yield"),
                ("stress_other_income","Other verified income"),("stress_distributions","Distributions paid"),
                ("stress_withdrawals","Withdrawals paid"),("stress_operations","Operations"),
            ]
            cols=st.columns(2)
            for i,(key,label) in enumerate(stress_labels):
                defaults[key]=cols[i%2].number_input(label,min_value=0.0,value=float(dec(settings[key])),step=100.0,key="stress_"+key)
            run_stress=st.form_submit_button("Run stress-test scenario",type="primary")
        if run_stress:
            save_settings(db_exec,record_event,user_email,{**defaults})
            st.session_state["cedars_stress_run"]=defaults
        data=st.session_state.get("cedars_stress_run",{k:float(dec(settings[k])) for k,_ in stress_labels})
        cash_in=sum(dec(data[k]) for k in ("stress_grid_revenue","stress_p2p_fees","stress_withdrawal_fees","stress_compounding","stress_staking","stress_other_income"))
        cash_out=sum(dec(data[k]) for k in ("stress_distributions","stress_withdrawals","stress_operations"))
        total=dec(data["stress_opening"])+cash_in-cash_out
        c1,c2,c3=st.columns(3)
        c1.metric("Scenario inflows",f"{float(cash_in):,.2f}")
        c2.metric("Scenario outflows",f"{float(cash_out):,.2f}")
        c3.metric("Ending scenario treasury",f"{float(total):,.2f}")
        monthly=[]
        balance=dec(data["stress_opening"])
        for month in range(1,25):
            inflow=cash_in/24; outflow=cash_out/24; opening=balance; balance=opening+inflow-outflow
            monthly.append({"Month":month,"Opening treasury (scenario)":float(opening),"New deposits":0.0,
                            "Scenario inflows":float(inflow),"Scenario outflows":float(outflow),
                            "Closing treasury (scenario)":float(balance),"Status":"SCENARIO ONLY"})
        st.dataframe(pd.DataFrame(monthly),use_container_width=True,hide_index=True)
        st.caption("Treasury yield inputs do not represent actual yield; they must not be counted as assets without evidence of receipt.")

    with tabs[12]:
        st.subheader("Executive dashboard — actual entries versus model counts")
        balances={}
        for item in ledger_rows:
            cur=str(item.get("currency") or "Other"); amt=dec(item.get("amount")); kind=str(item.get("entry_type") or "")
            if kind=="Verified income received": balances[cur]=balances.get(cur,Decimal(0))+amt
            elif kind in ("Verified payment sent","Verified expense","Refund"): balances[cur]=balances.get(cur,Decimal(0))-amt
        c1,c2,c3,c4=st.columns(4)
        c1.metric("Model-register investors",len(investor_rows))
        c2.metric("Grid evidence rows",len(revenue_rows))
        c3.metric("Ledger entries",len(ledger_rows))
        c4.metric("Currencies in ledger",len(balances))
        if balances:
            st.dataframe(pd.DataFrame([{"Currency":k,"Admin ledger net":float(v),"Source":"Internal ledger only; not live-verified"} for k,v in sorted(balances.items())]),use_container_width=True,hide_index=True)
        else: st.info("No treasury ledger movements have been recorded.")
        st.write("Pending withdrawals:",len(db_all("SELECT id FROM withdrawal_requests WHERE status NOT IN ('Confirmed on-chain','Rejected','Failed','Cancelled')")))
        st.write("Daily distribution target:",f"{float(dec(settings['daily_distribution'])):.2%} (model assumption only)")
        st.write("Reserve target:",f"{float(dec(settings['reserve_lock'])):.2%} (modeled; not actually ring-fenced)")
        st.info("No revenue, reserve, available balance, yield, or investor liability is asserted from this model unless supported by separate reconciled records.")

    solvency={key:settings.get(key,DEFAULTS[key]) for key in ("solvency_currency","solvency_assets","solvency_customer_balances","solvency_pending_rewards","solvency_pending_withdrawals")}
    _, health, action = solvency_status(solvency["solvency_assets"],solvency["solvency_customer_balances"],
        solvency["solvency_pending_rewards"],solvency["solvency_pending_withdrawals"],settings["solvency_threshold"],settings["healthy_threshold"])
    solvency.update({"assets":settings["solvency_assets"],"customer_balances":settings["solvency_customer_balances"],
        "pending_rewards":settings["solvency_pending_rewards"],"pending_withdrawals":settings["solvency_pending_withdrawals"],
        "currency":settings.get("solvency_currency","USDT"),"health":health,"action":action})
    stress={k:settings.get(k,DEFAULTS[k]) for k in DEFAULTS if k.startswith("stress_")}
    workbook_bytes=_wb_bytes(settings,investor_rows,revenue_rows,ledger_rows,solvency,stress)
    st.download_button("Download current 13-sheet Treasury Management Model (.xlsx)",data=workbook_bytes,
        file_name="cedars_empire_treasury_management_model_v1.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        help="Exports a snapshot of current model records; it does not enable or execute payments.")
