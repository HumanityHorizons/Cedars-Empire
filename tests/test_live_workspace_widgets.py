from io import BytesIO
from pathlib import Path
import sys

import pandas as pd
import pymupdf
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from hh_workspace_widgets import (
    GRID_COLUMNS, PAGE_ACTIONS, base_rows, _docx_bytes, _pdf_bytes
)


def test_every_live_page_has_a_configured_action_grid():
    assert PAGE_ACTIONS
    for page, actions in PAGE_ACTIONS.items():
        rows = base_rows(page)
        assert actions and len(rows) == len(actions)
        assert list(rows[0].keys()) == GRID_COLUMNS
        assert all(row["Action / requirement"].strip() for row in rows)


def test_word_export_is_readable_and_uses_one_point_five_spacing():
    frame = pd.DataFrame(base_rows("Opportunity Analysis"), columns=GRID_COLUMNS)
    doc = Document(BytesIO(_docx_bytes("Opportunity Analysis Grid", frame)))
    assert "Opportunity Analysis Grid" in doc.paragraphs[0].text
    assert all(p.paragraph_format.line_spacing == 1.5 for p in doc.paragraphs if p.text)


def test_pdf_export_is_valid_a4_and_contains_grid_content():
    frame = pd.DataFrame(base_rows("Opportunity Analysis"), columns=GRID_COLUMNS)
    pdf = pymupdf.open(stream=_pdf_bytes("Opportunity Analysis Grid", frame), filetype="pdf")
    assert pdf.page_count >= 1
    assert abs(pdf[0].rect.width - 595) < 1
    assert abs(pdf[0].rect.height - 842) < 1
    assert "Opportunity Analysis Grid" in pdf[0].get_text()
    pdf.close()
